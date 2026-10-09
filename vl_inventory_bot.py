#!/usr/bin/env python3
"""
Victory Live inventory Telegram bot.

Admins:
  /sync      — pull ticket groups from Victory Live into the bot database
  /autofulfill — fulfill all open VL mobile orders (barcodes from DB + stock files)
  /summary   — database counts
  /listed    — txt + Excel of listed and unlisted
  /bydate    — txt + Excel of inventory filtered by event date
  /stockfrom — non-SF tickets by event date range (ALL / barcode / app / no-info filters)
  /stats     — database totals and a live Victory Live count
  /list      — consign a listing
  /balance   — your balance

Super admin:
  /whitelist add|del|<blank to list>
  /blacklist add|del|<blank to list>
  /addbal <telegram_id> <amount> [reason]
"""
from __future__ import annotations

import asyncio
import calendar
import html
import json
import logging
import os
import re
import tempfile
import threading
from datetime import datetime
from pathlib import Path

from telegram import (
    BotCommand,
    BotCommandScopeAllPrivateChats,
    BotCommandScopeChat,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    Update,
)
from telegram.constants import ParseMode
from telegram.ext import (
    Application,
    ApplicationHandlerStop,
    CallbackQueryHandler,
    CommandHandler,
    ContextTypes,
    MessageHandler,
    filters,
)

ROOT = Path(__file__).resolve().parent
CFG_PATH = Path(os.environ.get("VL_BOT_CONFIG") or ROOT / "vl_inventory_bot.config.json")

# Hardcoded defaults when vl_inventory_bot.config.json is absent on VPS
DEFAULT_BOT_TOKEN = "8755470321:AAG1d95vryJHLnFGp4kH_DdoF_0-asGAXCg"
DEFAULT_ADMIN_IDS = "7550227847"


def _load_cfg() -> dict:
    if CFG_PATH.is_file():
        return json.loads(CFG_PATH.read_text(encoding="utf-8"))
    return {}


CFG = _load_cfg()

from vl_auth import apply_proxy_settings

apply_proxy_settings(
    use_proxy=bool(CFG.get("use_proxy", True)),
    proxy_url=(str(CFG.get("vl_proxy_url") or "").strip() or None),
    proxy_user=(str(CFG.get("vl_proxy_user") or "").strip() or None),
    proxy_pass=(str(CFG.get("vl_proxy_pass") or "").strip() or None),
    cookies_only=bool(CFG.get("vl_cookies_only", True)),
    vl_username=(str(CFG.get("vl_username") or "").strip() or None),
    vl_password=CFG.get("vl_password") if CFG.get("vl_password") is not None else None,
)

import vl_inventory_db as db
import vl_stock
from vl_stock import event_date_to_iso
from vl_inventory_sync import fix_all_pos_stock_types, live_inventory_stats, sync_inventory, sync_orders
from vl_listing_handlers import build_listing_conversation
from vl_listing_user_report import build_ibrahim_report, format_ibrahim_report_html

BOT_TOKEN = os.environ.get("VL_BOT_TOKEN") or CFG.get("bot_token") or DEFAULT_BOT_TOKEN
ADMIN_IDS = {
    int(x)
    for x in str(os.environ.get("VL_ADMIN_IDS") or CFG.get("admin_ids") or DEFAULT_ADMIN_IDS).split(",")
    if str(x).strip().isdigit()
}
CHAT_IDS = {
    int(x)
    for x in str(os.environ.get("VL_NOTIFY_CHAT_IDS") or CFG.get("notify_chat_ids") or "").split(",")
    if str(x).strip().isdigit()
} | ADMIN_IDS
AFFILIATE_GROUP_ID = int(CFG.get("affiliate_group_id") or 722)
SYNC_DAYS_BACK = int(CFG.get("sync_days_back") or 7)
SYNC_DAYS_FORWARD = int(CFG.get("sync_days_forward") or 365)
# True = all VL HTTP via vl_proxy_url (fixed egress on 96.126.176.73:8749), not direct VPS IP
USE_PROXY = bool(CFG.get("use_proxy", True))
SUPER_ADMIN_ID = int(CFG.get("super_admin_id") or 7550227847)
CLIENT_PO_PROFILE_ID = int(CFG.get("client_po_profile_id") or 80001)
CONSIGN_PROFILE_NAME = str(CFG.get("consign_profile_name") or "Consign - Fouad Mahmoud")
LISTING_PAYOUT_USD = float(CFG.get("listing_payout_usd") or 0.25)

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
log = logging.getLogger("vl_inventory_bot")


def _admin(uid: int) -> bool:
    return uid in ADMIN_IDS


def _worker(uid: int) -> bool:
    return db.is_worker(uid)


def _can_list(uid: int) -> bool:
    if db.is_blacklisted(uid) and not _admin(uid):
        return False
    return _admin(uid) or _worker(uid)


def _super_admin(uid: int) -> bool:
    return uid == SUPER_ADMIN_ID


_WORKER_ONLY_HINT = "That command is for admins only. You can use /list and /balance."


async def _deny_unless_admin(update: Update) -> bool:
    """Return True if the user is an admin. Otherwise reply and return False."""
    if not update.effective_user or not update.message:
        return False
    uid = update.effective_user.id
    if _admin(uid):
        return True
    if _worker(uid):
        await update.message.reply_text(_WORKER_ONLY_HINT)
    else:
        await update.message.reply_text("Unauthorized.")
    return False


def _seats(row) -> str:
    start = str(row["start_seat"] or "")
    end = str(row["end_seat"] or "")
    if end and end != start:
        return f"{start}-{end}"
    return start


async def start_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not update.effective_user or not update.message:
        return
    uid = update.effective_user.id
    if db.is_blacklisted(uid) and not _admin(uid):
        await update.message.reply_text("You are blacklisted.")
        return
    if _can_list(uid) and not _admin(uid):
        await update.message.reply_text(
            "Victory Live + TM worker bot.\n"
            "/list — consign on Victory Live\n"
            "/work — TM account listing (barcodes / LISTED)\n"
            "/status — current TM job\n"
            "/cancel — exit TM step\n"
            "/leaderboard — top workers\n"
            "/balance — VL + TM earnings"
        )
        return
    if not _admin(uid):
        await update.message.reply_text("Unauthorized.")
        return
    text = (
        "Victory Live bot.\n"
        "/sync — inventory and orders, files sent when it finishes\n"
        "/autofulfill — fulfill open VL mobile orders using DB barcodes\n"
        "/cancel_autofulfill — stop a running auto-fulfill and send partial report\n"
        "/autofulfill_reset — clear tried history so next run hits all orders again\n"
        "/fulfill_file — upload checker barcode file for autofulfill (merges)\n"
        "/fixstocktype — set all VL stock types to MOBILE TRANSFER (use dry to preview)\n"
        "/summary — database counts\n"
        "/listed — txt + Excel of listed and unlisted\n"
        "/bydate <date> [to] — txt + Excel of inventory filtered by event date\n"
        "/stats — database totals and live Victory Live totals\n"
        "/fetchibrahim — Ibrahim lister DB + VL totals (admin)\n"
        "/list — consign a listing\n"
        "/balance — your balance\n"
        "/stock — ticket database totals\n"
        "/stockfind <text> — search stored tickets\n"
        "/stockfile — txt + Excel of every stored ticket\n"
        "/stockfrom — non-SF export: start + end event dates, then filter buttons\n"
        "Send a .txt, .csv, or .xlsx to import ticket DB (app/barcode/seats)\n\n"
        "Databases:\n"
        "/dbstatus — VL stock + TM batches + SF split summary\n"
        "/stockdedupe — remove duplicate rows in ticket DB (keeps newest)\n\n"
        "TM worker (/work) — data in data/tm_worker/:\n"
        "/upload then file, or auto-detect TM .xlsx/.txt\n"
        "/admin — full TM admin command list\n"
        "/batches /queue /workers /stats /collect /export_all /price /dashboard …\n"
        "Workers: /work /status /cancel /leaderboard"
    )
    if _super_admin(uid):
        text += (
            "\n\nAdmin:\n"
            "/whitelist add <id> [note]\n"
            "/whitelist del <id>\n"
            "/whitelist — list users and balances\n"
            "/blacklist add <id> [note]\n"
            "/blacklist del <id>\n"
            "/blacklist — list blocked users\n"
            "/addbal <id> <amount> [reason]"
        )
    await update.message.reply_text(text)


def _parse_user_id(token: str) -> int | None:
    try:
        return int(token)
    except (TypeError, ValueError):
        return None


async def whitelist_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not update.effective_user or not update.message or not _super_admin(update.effective_user.id):
        if update.message:
            await update.message.reply_text("Only the super admin can manage the whitelist.")
        return
    args = list(context.args or [])
    action = (args[0].lower() if args else "list")
    if action in ("list", "ls"):
        rows = db.list_workers()
        if not rows:
            await update.message.reply_text("Whitelist is empty.")
            return
        lines = ["<b>Whitelist</b>"]
        for r in rows:
            bal = db.get_balance(int(r["telegram_id"]))
            note = f" — {html.escape(r['note'])}" if r["note"] else ""
            lines.append(
                f"<code>{r['telegram_id']}</code> ${bal['balance_usd']:.2f} "
                f"({bal['listings_paid']} paid){note}"
            )
        await update.message.reply_text("\n".join(lines), parse_mode=ParseMode.HTML)
        return
    if action not in ("add", "del", "remove") or len(args) < 2:
        await update.message.reply_text("Usage: /whitelist add <id> [note] | /whitelist del <id> | /whitelist")
        return
    wid = _parse_user_id(args[1])
    if wid is None:
        await update.message.reply_text("Invalid telegram id.")
        return
    if action == "add":
        db.remove_blacklist(wid)
        note = " ".join(args[2:]).strip()
        added = db.add_worker(wid, update.effective_user.id, note)
        try:
            from ticketmaster.tm_worker_bridge import sync_worker_to_tm_whitelist

            sync_worker_to_tm_whitelist(wid)
        except Exception:
            log.exception("TM whitelist sync failed for %s", wid)
        await update.message.reply_text(
            f"Whitelisted {wid}." if added else f"{wid} was already whitelisted."
        )
        return
    if db.remove_worker(wid):
        try:
            from ticketmaster.tm_worker_bridge import remove_worker_from_tm_whitelist

            remove_worker_from_tm_whitelist(wid)
        except Exception:
            log.exception("TM whitelist remove failed for %s", wid)
        await update.message.reply_text(f"Removed {wid} from the whitelist.")
    else:
        await update.message.reply_text(f"{wid} was not on the whitelist.")


async def blacklist_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not update.effective_user or not update.message or not _super_admin(update.effective_user.id):
        if update.message:
            await update.message.reply_text("Only the super admin can manage the blacklist.")
        return
    args = list(context.args or [])
    action = (args[0].lower() if args else "list")
    if action in ("list", "ls"):
        rows = db.list_blacklist()
        if not rows:
            await update.message.reply_text("Blacklist is empty.")
            return
        lines = ["<b>Blacklist</b>"]
        for r in rows:
            note = f" — {html.escape(r['note'])}" if r["note"] else ""
            lines.append(f"<code>{r['telegram_id']}</code>{note}")
        await update.message.reply_text("\n".join(lines), parse_mode=ParseMode.HTML)
        return
    if action not in ("add", "del", "remove") or len(args) < 2:
        await update.message.reply_text("Usage: /blacklist add <id> [note] | /blacklist del <id> | /blacklist")
        return
    wid = _parse_user_id(args[1])
    if wid is None:
        await update.message.reply_text("Invalid telegram id.")
        return
    if action == "add":
        note = " ".join(args[2:]).strip()
        added = db.add_blacklist(wid, update.effective_user.id, note)
        await update.message.reply_text(
            f"Blacklisted {wid}. They cannot /list." if added else f"{wid} was already blacklisted."
        )
        return
    if db.remove_blacklist(wid):
        await update.message.reply_text(f"Removed {wid} from the blacklist.")
    else:
        await update.message.reply_text(f"{wid} was not blacklisted.")


async def addbal_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not update.effective_user or not update.message or not _super_admin(update.effective_user.id):
        if update.message:
            await update.message.reply_text("Only the super admin can change balances.")
        return
    args = list(context.args or [])
    if len(args) < 2:
        await update.message.reply_text("Usage: /addbal <telegram_id> <amount> [reason]")
        return
    wid = _parse_user_id(args[0])
    if wid is None:
        await update.message.reply_text("Invalid telegram id.")
        return
    try:
        amount = float(args[1].replace(",", "").replace("$", ""))
    except ValueError:
        await update.message.reply_text("Amount must be a number. Use a negative amount to remove balance.")
        return
    reason = " ".join(args[2:]).strip()
    result = db.adjust_balance(wid, amount, added_by=update.effective_user.id, reason=reason)
    sign = "+" if amount >= 0 else ""
    await update.message.reply_text(
        f"{sign}${amount:.2f} for <code>{wid}</code>\n"
        f"Balance: <b>${result['balance_usd']:.2f}</b>",
        parse_mode=ParseMode.HTML,
    )


async def balance_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not update.effective_user or not update.message:
        return
    uid = update.effective_user.id
    if not _can_list(uid):
        await update.message.reply_text("Unauthorized.")
        return
    bal = db.get_balance(uid)
    lines = [
        f"<b>Victory Live</b> balance: <b>${bal['balance_usd']:.2f}</b>",
        f"VL listings paid: {bal['listings_paid']}",
    ]
    try:
        from ticketmaster import bot as tm

        tm.db.ensure_worker(uid, update.effective_user.username, update.effective_user.full_name)
        tb = tm.db.get_worker_balance(uid)
        rate = tm.db.get_payout_rate_for_user(uid)
        lines.append(
            f"\n<b>TM / StageFront worker</b> balance: <b>${tb['balance_usd']:.2f}</b>"
        )
        lines.append(
            f"TM listings this period: {tb.get('listings_completed', tb['barcodes_submitted'])} "
            f"(pay ${rate:.2f}/listing when marked LISTED)"
        )
    except Exception:
        pass
    await update.message.reply_text("\n".join(lines), parse_mode=ParseMode.HTML)


_SYNC_LOCK = asyncio.Lock()


def _write_sync_files() -> dict[str, Path]:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    folder = Path(tempfile.mkdtemp(prefix="vl_sync_"))
    groups = db.export_active_groups()
    orders = db.list_orders()
    inv_headers = [
        "ticket_group_id",
        "event",
        "date",
        "venue",
        "section",
        "row",
        "seats",
        "qty",
        "list",
        "cost",
        "status",
        "state",
        "notes",
    ]
    inv_rows = []
    for row in groups:
        notes = " | ".join(
            part for part in (row["internal_notes"] or "", row["external_notes"] or "") if part
        )
        inv_rows.append(
            {
                "ticket_group_id": row["ticket_group_id"],
                "event": row["production_name"] or "",
                "date": row["production_date"] or "",
                "venue": row["venue"] or "",
                "section": row["section"] or "",
                "row": row["row_name"] or "",
                "seats": _seats(row),
                "qty": int(row["quantity"] or 0),
                "list": float(row["list_price"] or 0),
                "cost": float(row["cost"] or 0),
                "status": "listed" if row["is_shared"] else "unlisted",
                "state": row["state"] or "",
                "notes": notes,
            }
        )
    inv_json = folder / f"inventory_{stamp}.json"
    inv_txt = folder / f"inventory_{stamp}.txt"
    inv_json.write_text(json.dumps(inv_rows, indent=2), encoding="utf-8")
    inv_lines = ["\t".join(inv_headers)]
    inv_lines.extend("\t".join(str(item[h]).replace("\t", " ") for h in inv_headers) for item in inv_rows)
    inv_txt.write_text("\n".join(inv_lines) + "\n", encoding="utf-8")

    order_rows = []
    for row in orders:
        try:
            tickets = json.loads(row["ticket_transfer_urls"] or "[]")
        except json.JSONDecodeError:
            tickets = []
        order_rows.append(
            {
                "rex_order_id": row["rex_order_id"],
                "invoice": row["invoice_number"] or "",
                "bucket": row["bucket"] or "",
                "status": row["order_status"] or "",
                "internal_status": row["internal_status"] or "",
                "marketplace": row["marketplace"] or "",
                "marketplace_order_id": row["marketplace_order_id"] or "",
                "event": row["production_name"] or "",
                "date": row["production_date"] or "",
                "venue": row["venue"] or "",
                "section": row["section"] or "",
                "row": row["row_name"] or "",
                "seats": row["seats"] or "",
                "qty": row["quantity"] or 0,
                "total": row["total_cost"],
                "payout": row["payout"],
                "fulfillment": row["fulfillment_type"] or "",
                "buyer_email": row["buyer_email"] or "",
                "customer_first": row["customer_first"] or "",
                "customer_last": row["customer_last"] or "",
                "customer_email": row["customer_email"] or "",
                "customer_account_email": row["customer_account_email"] or "",
                "transfer_url": row["transfer_url"] or "",
                "transfer_confirmation": row["transfer_confirmation"] or "",
                "tickets": tickets,
                "status_message": row["status_message"] or "",
                "created": row["create_date"] or "",
            }
        )
    ord_json = folder / f"orders_{stamp}.json"
    ord_txt = folder / f"orders_{stamp}.txt"
    ord_json.write_text(json.dumps(order_rows, indent=2), encoding="utf-8")
    txt_lines = [f"Orders {stamp}", ""]
    for item in order_rows:
        txt_lines.extend(
            [
                f"{item['bucket'].upper()}  Inv {item['invoice'] or '—'}  {item['status']}  TG-order {item['rex_order_id']}",
                f"{item['event']} : {item['venue']} : {item['date']}",
                f"Sec {item['section']} Row {item['row']} Seats {item['seats']} qty {item['qty']}  {item['marketplace']} {item['marketplace_order_id']}",
                f"Customer: {_customer_line_dict(item) or '—'}",
                f"Buyer: {item['buyer_email'] or '—'}",
                f"Transfer URL: {item['transfer_url'] or '—'}",
                f"Confirmation #: {item['transfer_confirmation'] or '—'}",
            ]
        )
        for ticket in item["tickets"]:
            if ticket.get("transfer_url"):
                txt_lines.append(
                    f"  Seat {ticket.get('seat')} {ticket.get('section')}/{ticket.get('row')} {ticket.get('transfer_url')}"
                )
        if item["status_message"]:
            txt_lines.append(item["status_message"])
        txt_lines.append("")
    ord_txt.write_text("\n".join(txt_lines), encoding="utf-8")
    return {
        "inventory_json": inv_json,
        "inventory_txt": inv_txt,
        "orders_json": ord_json,
        "orders_txt": ord_txt,
        "groups": len(inv_rows),
        "orders": len(order_rows),
    }


def _customer_line_dict(item: dict) -> str:
    name = " ".join(part for part in (item.get("customer_first") or "", item.get("customer_last") or "") if part).strip()
    email = item.get("customer_email") or item.get("buyer_email") or ""
    if name and email:
        return f"{name} <{email}>"
    return name or email


def _sync_result_text(inventory: dict | None, orders: dict | None) -> str:
    lines = ["Sync finished"]
    if inventory and inventory.get("ok"):
        snap = db.summary()
        lines.extend(
            [
                "",
                "Inventory",
                f"Productions: {inventory['productions']} · Groups: {inventory['groups']}",
                f"Window: {inventory['from']} → {inventory['to']}",
                f"Active: {snap['active']} · Listed: {snap['shared']} · Unlisted: {snap['active_unshared']}",
                f"List value: {_money(snap['active_list_value'])} · Cost value: {_money(snap['active_cost_value'])}",
            ]
        )
    elif inventory:
        lines.append(f"Inventory failed: {inventory.get('error')}")
    if orders and orders.get("ok"):
        lines.extend(
            [
                "",
                "Orders",
                f"Created: {orders['from']} → {orders['to']}",
                f"Total: {orders['orders']}",
                f"Pending: {orders['pending']}",
                f"Completed: {orders['completed']}",
                f"Canceled: {orders['canceled']}",
                f"With transfer URL: {orders['with_transfer']}",
                f"With confirmation #: {orders.get('with_confirmation', 0)}",
                f"With customer details: {orders['with_customer']}",
            ]
        )
    elif orders:
        lines.append(f"Orders failed: {orders.get('error')}")
    return "\n".join(lines)


async def _run_sync(bot, chat_id: int, message_id: int) -> None:
    async with _SYNC_LOCK:
        inventory = None
        orders = None
        try:
            await bot.edit_message_text(
                "Sync running. Pulling current inventory. Other commands still work.",
                chat_id=chat_id,
                message_id=message_id,
            )
            inventory = await asyncio.to_thread(
                sync_inventory,
                days_back=SYNC_DAYS_BACK,
                days_forward=SYNC_DAYS_FORWARD,
                affiliate_group_id=AFFILIATE_GROUP_ID,
                use_proxy=USE_PROXY,
            )
            await bot.edit_message_text(
                "Inventory saved. Pulling orders, transfer URLs, and customer details.",
                chat_id=chat_id,
                message_id=message_id,
            )
            try:
                orders = await asyncio.to_thread(sync_orders, use_proxy=USE_PROXY)
            except Exception as exc:
                orders = {"ok": False, "error": str(exc)}
            files = await asyncio.to_thread(_write_sync_files)
            caption = _sync_result_text(inventory, orders)
            await bot.edit_message_text(caption[:4000], chat_id=chat_id, message_id=message_id)
            for key in ("inventory_txt", "inventory_json", "orders_txt", "orders_json"):
                path = files[key]
                with path.open("rb") as handle:
                    await bot.send_document(
                        chat_id=chat_id,
                        document=handle,
                        filename=path.name,
                        caption=f"{files['groups']} groups · {files['orders']} orders",
                    )
        except Exception as exc:
            log.exception("sync failed")
            await bot.edit_message_text(
                f"Sync failed: {exc}\n\n{_sync_result_text(inventory, orders)}"[:4000],
                chat_id=chat_id,
                message_id=message_id,
            )
        finally:
            pass


_AUTOFULFILL_LOCK = asyncio.Lock()
_AUTOFULFILL_CANCEL_EVENT: threading.Event | None = None


def _autofulfill_fail_bucket(message: str) -> str:
    m = (message or "").strip().lower()
    if not m:
        return "unknown"
    if "no barcode" in m:
        return "no barcode in DB"
    if "purchase manag" in m or "season seat" in m or "barcodes detected" in m:
        return "VL purchase manager / season seat"
    if "not matched" in m or "unable to transfer" in m:
        return "transfer rejected by TM/VL"
    if m.startswith("update failed") or "barcodes still empty" in m:
        return "barcode update failed"
    if m.startswith("fulfill http") or "fulfill rejected" in m:
        return "fulfill API error"
    if "session expired" in m:
        return "session expired"
    if m == "no tickets":
        return "no mobile tickets"
    return "other"


def _write_autofulfill_report_files(result: dict) -> tuple[Path, Path]:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    folder = Path(tempfile.mkdtemp(prefix="vl_autofulfill_"))
    txt_path = folder / f"autofulfill_report_{stamp}.txt"
    json_path = folder / f"autofulfill_report_{stamp}.json"

    results = list(result.get("results") or [])
    ok_rows = [r for r in results if r.get("success")]
    fail_rows = [r for r in results if not r.get("success")]
    buckets: dict[str, int] = {}
    for r in fail_rows:
        b = _autofulfill_fail_bucket(str(r.get("message") or ""))
        buckets[b] = buckets.get(b, 0) + 1

    def _sort_key(r: dict) -> tuple:
        return (str(r.get("invoice") or ""),)

    ok_rows.sort(key=_sort_key)
    fail_rows.sort(key=_sort_key)

    elapsed = result.get("elapsed_sec")
    lines = [
        "Victory Live auto-fulfill report",
        f"Generated: {datetime.now().isoformat(timespec='seconds')}",
        "",
        "Summary",
        f"  Duration: {elapsed}s" if elapsed is not None else "  Duration: —",
        f"  Workers: {result.get('workers', '?')}",
        f"  Stock rows (barcodes): {result.get('stock_rows', '?')}",
        f"  Open VL orders scanned: {result.get('open_orders', '?')}",
        f"  Mobile unfulfilled targets: {result.get('targets', '?')}",
        f"  Mobile candidates: {result.get('mobile_candidates', '—')}",
        f"  Merged from local DB: {result.get('orders_merged_from_db', 0)}",
        f"  New this run: {result.get('targets_new', result.get('targets', '?'))}",
        f"  Skipped (already tried): {result.get('skipped_tried', 0)}",
        f"  On file total tried: {result.get('tried_total', '?')}",
        f"  Success: {result.get('ok', 0)}",
        f"  Failed: {result.get('fail', 0)}",
    ]
    if result.get("cancelled"):
        lines.extend(
            [
                f"  Cancelled: yes ({result.get('targets_remaining', 0)} not attempted)",
                "",
            ]
        )
    else:
        lines.append("")
    if buckets:
        lines.append("Failure breakdown")
        for name, count in sorted(buckets.items(), key=lambda x: (-x[1], x[0])):
            lines.append(f"  {count:4d}  {name}")
        lines.append("")

    lines.append(f"=== SUCCESS ({len(ok_rows)}) ===")
    if ok_rows:
        for r in ok_rows:
            inv = r.get("invoice") or "?"
            ev = (r.get("event") or "")[:70]
            msg = str(r.get("message") or "").replace("\n", " ")[:200]
            lines.append(f"OK  {inv}  |  {ev}  |  {msg}")
    else:
        lines.append("(none)")
    lines.append("")
    lines.append(f"=== FAILED ({len(fail_rows)}) ===")
    if fail_rows:
        for r in fail_rows:
            inv = r.get("invoice") or "?"
            ev = (r.get("event") or "")[:70]
            msg = str(r.get("message") or "").replace("\n", " ")[:300]
            lines.append(f"FAIL  {inv}  |  {ev}  |  {msg}")
    else:
        lines.append("(none)")

    not_attempted = list(result.get("not_attempted") or [])
    if not_attempted:
        lines.extend(
            [
                "",
                f"=== NOT ATTEMPTED ({len(not_attempted)}) — run cancelled or stopped early ===",
            ]
        )
        for row in not_attempted:
            inv = row.get("invoice") or "?"
            ev = (row.get("event") or "")[:70]
            sec = row.get("section") or ""
            rw = row.get("row") or ""
            seats = row.get("seats") or ""
            lines.append(f"SKIP  {inv}  |  {ev}  |  {sec} {rw} {seats}")

    txt_path.write_text("\n".join(lines), encoding="utf-8")
    json_path.write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")
    return txt_path, json_path


def _autofulfill_summary_html(result: dict) -> str:
    ok = int(result.get("ok") or 0)
    fail = int(result.get("fail") or 0)
    targets = int(result.get("targets") or 0)
    elapsed = result.get("elapsed_sec")
    workers = result.get("workers", "?")

    results = list(result.get("results") or [])
    fail_rows = [r for r in results if not r.get("success")]
    buckets: dict[str, int] = {}
    for r in fail_rows:
        b = _autofulfill_fail_bucket(str(r.get("message") or ""))
        buckets[b] = buckets.get(b, 0) + 1

    title = (
        "<b>Auto-fulfill cancelled</b> (partial report)"
        if result.get("cancelled")
        else "<b>Auto-fulfill finished</b>"
    )
    lines = [
        title,
        "",
        f"Duration: <b>{elapsed}s</b> · workers: <b>{workers}</b>" if elapsed else f"Workers: <b>{workers}</b>",
        f"Stock barcodes: <b>{result.get('stock_rows', '?')}</b>"
        + (
            f" (<b>{result.get('fulfill_file_rows', 0)}</b> from /fulfill_file)"
            if result.get("fulfill_file_rows")
            else ""
        ),
        f"Scanned open orders: <b>{result.get('open_orders', '?')}</b>",
        f"Mobile unfulfilled targets: <b>{targets}</b> · new this run: <b>{result.get('targets_new', ok + fail)}</b>",
        f"Success: <b>{ok}</b> · Failed: <b>{fail}</b>",
    ]
    if result.get("cancelled"):
        rem = int(result.get("targets_remaining") or 0)
        lines.append(f"Stopped early — <b>{rem}</b> target(s) not attempted (see report)")
    mc = result.get("mobile_candidates")
    if mc is not None:
        lines.append(f"Mobile candidates: <b>{mc}</b> (not all become fulfill targets)")
    if result.get("orders_merged_from_db"):
        lines.append(f"Added from local DB (/sync): <b>{result.get('orders_merged_from_db')}</b>")
    skipped = int(result.get("skipped_tried") or 0)
    if skipped:
        lines.append(f"Skipped (already tried): <b>{skipped}</b>")
    tried_total = result.get("tried_total")
    if tried_total is not None:
        lines.append(f"Tried file entries: <b>{tried_total}</b> (<code>data/vl_autofulfill_tried.json</code>)")
    run_n = int(result.get("targets_new") or ok + fail)
    if run_n:
        pct = int(round(100 * ok / run_n))
        lines.append(f"Success rate (this run): <b>{pct}%</b>")
    elif targets and skipped:
        lines.append("<i>No new orders to attempt.</i>")

    if buckets:
        lines.extend(["", "<b>Why failures</b>"])
        for name, count in sorted(buckets.items(), key=lambda x: (-x[1], x[0]))[:8]:
            lines.append(f"• {html.escape(name)}: <b>{count}</b>")

    ok_rows = [r for r in results if r.get("success")]
    if ok_rows:
        lines.extend(["", "<b>OK</b> (sample)"])
        for r in ok_rows[:8]:
            inv = html.escape(str(r.get("invoice") or "?"))
            ev = html.escape((r.get("event") or "")[:45])
            lines.append(f"✓ {inv} — {ev}")

    if fail_rows:
        lines.extend(["", "<b>FAIL</b> (sample)"])
        for r in fail_rows[:12]:
            inv = html.escape(str(r.get("invoice") or "?"))
            msg = html.escape(str(r.get("message") or "")[:100])
            lines.append(f"✗ {inv}: {msg}")
        if len(fail_rows) > 12:
            lines.append(f"… +{len(fail_rows) - 12} more in attached report")

    lines.append("")
    lines.append("Full <b>.txt</b> + <b>.json</b> report attached below.")
    text = "\n".join(lines)
    return text[:3900]


async def _send_autofulfill_report(
    bot,
    chat_id: int,
    *,
    summary_message_id: int,
    result: dict,
) -> None:
    txt_path, json_path = _write_autofulfill_report_files(result)
    try:
        await bot.edit_message_text(
            chat_id=chat_id,
            message_id=summary_message_id,
            text=_autofulfill_summary_html(result),
            parse_mode=ParseMode.HTML,
        )
    except Exception:
        pass
    cap = (
        f"Auto-fulfill · ok {result.get('ok', 0)} fail {result.get('fail', 0)} "
        f"· {result.get('targets', 0)} targets"
    )
    try:
        with txt_path.open("rb") as f:
            await bot.send_document(chat_id=chat_id, document=f, filename=txt_path.name, caption=cap)
        with json_path.open("rb") as f:
            await bot.send_document(chat_id=chat_id, document=f, filename=json_path.name)
    except Exception as exc:
        log.exception("autofulfill report upload failed: %s", exc)
        try:
            await bot.send_message(
                chat_id=chat_id,
                text=f"Report files failed to send: {exc}"[:500],
            )
        except Exception:
            pass
    finally:
        try:
            import shutil

            shutil.rmtree(txt_path.parent, ignore_errors=True)
        except Exception:
            pass


async def cancel_autofulfill_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not await _deny_unless_admin(update):
        return
    global _AUTOFULFILL_CANCEL_EVENT
    if not _AUTOFULFILL_LOCK.locked():
        await update.message.reply_text("No auto-fulfill is running.")
        return
    if _AUTOFULFILL_CANCEL_EVENT is None:
        await update.message.reply_text("Auto-fulfill is running but cancel is not available yet — try again in a few seconds.")
        return
    _AUTOFULFILL_CANCEL_EVENT.set()
    await update.message.reply_text(
        "Cancelling auto-fulfill…\n"
        "In-flight orders may still finish; then you get the partial <b>.txt</b> + <b>.json</b> report.",
        parse_mode=ParseMode.HTML,
    )


async def autofulfill_reset_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not await _deny_unless_admin(update):
        return
    if _AUTOFULFILL_LOCK.locked():
        await update.message.reply_text(
            "Auto-fulfill is running — use <code>/cancel_autofulfill</code> to stop and get a partial report.",
            parse_mode=ParseMode.HTML,
        )
        return
    from vl_autofulfill_tried import clear_tried_history, tried_count

    removed = await asyncio.to_thread(clear_tried_history)
    await update.message.reply_text(
        f"Cleared autofulfill history: <b>{removed}</b> tried order(s) removed.\n"
        f"Tried file now: <b>{await asyncio.to_thread(tried_count)}</b> entries.\n\n"
        "Next <code>/autofulfill</code> will attempt every open mobile target again.",
        parse_mode=ParseMode.HTML,
    )


async def autofulfill_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not await _deny_unless_admin(update):
        return
    if _AUTOFULFILL_LOCK.locked():
        await update.message.reply_text("Auto-fulfill is already running.")
        return

    msg = await update.message.reply_text(
        "Auto-fulfill started — scanning VL orders "
        "(created: 48mo · events: 24mo back + through Dec 2027 · local /sync pending). "
        "When done you get a summary plus txt/json report files."
    )
    chat_id = update.effective_chat.id
    mid = msg.message_id
    loop = asyncio.get_running_loop()
    q: asyncio.Queue[str | None] = asyncio.Queue()
    last_push = {"t": 0.0}

    def on_progress(text: str) -> None:
        import time as _time

        now = _time.monotonic()
        if now - last_push["t"] < 1.0 and not (
            text.startswith("Fulfilling") or text.startswith("FAIL ")
        ):
            return
        last_push["t"] = now
        try:
            loop.call_soon_threadsafe(q.put_nowait, f"Auto-fulfill: {text}"[:1000])
        except Exception:
            pass

    async def _pump() -> None:
        last = ""
        while True:
            text = await q.get()
            if text is None:
                return
            if text == last:
                continue
            last = text
            try:
                await context.bot.edit_message_text(chat_id=chat_id, message_id=mid, text=text)
            except Exception:
                pass

    async def _run() -> None:
        global _AUTOFULFILL_CANCEL_EVENT
        cancel_event = threading.Event()
        _AUTOFULFILL_CANCEL_EVENT = cancel_event
        pump = asyncio.create_task(_pump())
        try:
            async with _AUTOFULFILL_LOCK:
                try:
                    from ticketmaster.vl_fulfill_service import run_autofulfill_all_unfulfilled

                    result = await asyncio.to_thread(
                        run_autofulfill_all_unfulfilled,
                        on_progress=on_progress,
                        cancel_event=cancel_event,
                    )
                except Exception as exc:
                    log.exception("autofulfill failed")
                    await q.put(None)
                    await pump
                    try:
                        await context.bot.edit_message_text(
                            chat_id=chat_id,
                            message_id=mid,
                            text=f"Auto-fulfill failed: {exc}"[:4000],
                        )
                    except Exception:
                        pass
                    return
                await q.put(None)
                await pump
                await _send_autofulfill_report(
                    context.bot,
                    chat_id,
                    summary_message_id=mid,
                    result=result,
                )
        finally:
            _AUTOFULFILL_CANCEL_EVENT = None

    asyncio.create_task(_run())


async def sync_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not await _deny_unless_admin(update):
        return
    if _SYNC_LOCK.locked():
        await update.message.reply_text("Sync is already running. Other commands still work.")
        return
    msg = await update.message.reply_text(
        "Sync started. It will send the inventory and orders as JSON and txt when it finishes. Other commands still work."
    )
    asyncio.create_task(_run_sync(context.bot, update.effective_chat.id, msg.message_id))


async def fixstocktype_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not await _deny_unless_admin(update):
        return
    if _SYNC_LOCK.locked():
        await update.message.reply_text("Sync is already running — wait or try again in a minute.")
        return
    dry = bool(context.args and context.args[0].lower() in ("dry", "dry-run", "preview"))
    msg = await update.message.reply_text(
        "Scanning Victory Live inventory and setting stock type to <b>MOBILE TRANSFER</b>…"
        + (" (dry run)" if dry else ""),
        parse_mode=ParseMode.HTML,
    )

    async def _run() -> None:
        async with _SYNC_LOCK:
            try:
                result = await asyncio.to_thread(
                    fix_all_pos_stock_types,
                    days_back=SYNC_DAYS_BACK,
                    days_forward=SYNC_DAYS_FORWARD,
                    affiliate_group_id=AFFILIATE_GROUP_ID,
                    use_proxy=USE_PROXY,
                    dry_run=dry,
                )
            except Exception as exc:
                log.exception("fixstocktype failed")
                await context.bot.edit_message_text(
                    chat_id=update.effective_chat.id,
                    message_id=msg.message_id,
                    text=f"Fix stock type failed: {exc}"[:4000],
                )
                return
            if not result.get("ok"):
                await context.bot.edit_message_text(
                    chat_id=update.effective_chat.id,
                    message_id=msg.message_id,
                    text=f"Fix stock type failed: {result.get('error', result)}"[:4000],
                )
                return
            if dry:
                text = (
                    f"Dry run · scanned {result.get('scanned', 0)} groups\n"
                    f"Would update: {result.get('would_fix', 0)}"
                )
            else:
                text = (
                    f"Done · scanned {result.get('scanned', 0)} groups\n"
                    f"Updated stock type on {result.get('fixed', 0)} groups to MOBILE TRANSFER"
                )
            await context.bot.edit_message_text(
                chat_id=update.effective_chat.id,
                message_id=msg.message_id,
                text=text,
            )

    asyncio.create_task(_run())


def _money(n: float) -> str:
    return f"${n:,.2f}"


def _db_stats_text(s: dict) -> str:
    return (
        "Database\n"
        f"Groups: {s['total']}\n"
        f"Active: {s['active']} · tickets {s['active_qty']}\n"
        f"Listed: {s['shared']} · tickets {s['listed_qty']} · list {_money(s['listed_list_value'])}\n"
        f"Unlisted: {s['active_unshared']}\n"
        f"Active list value: {_money(s['active_list_value'])}\n"
        f"Active cost value: {_money(s['active_cost_value'])}\n"
        f"Last sync: {s['last_sync'] or 'never'} ({s['last_sync_groups']} groups)"
    )


async def summary_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not await _deny_unless_admin(update):
        return
    await update.message.reply_text(_db_stats_text(db.summary()))


_EXPORT_HEADERS = [
    "ticket_group_id",
    "event",
    "date",
    "venue",
    "section",
    "row",
    "seats",
    "qty",
    "list",
    "cost",
    "status",
    "notes",
]


def _export_values(row) -> list:
    notes = " | ".join(
        part for part in (row["internal_notes"] or "", row["external_notes"] or "") if part
    )
    return [
        row["ticket_group_id"],
        row["production_name"] or "",
        row["production_date"] or "",
        row["venue"] or "",
        row["section"] or "",
        row["row_name"] or "",
        _seats(row),
        int(row["quantity"] or 0),
        float(row["list_price"] or 0),
        float(row["cost"] or 0),
        "listed" if row["is_shared"] else "unlisted",
        notes,
    ]


def _split_listed(rows: list) -> tuple[list, list]:
    listed = [r for r in rows if r["is_shared"]]
    unlisted = [r for r in rows if not r["is_shared"]]
    return listed, unlisted


def _write_inventory_txt(rows: list) -> tuple[Path, int, int]:
    listed, unlisted = _split_listed(rows)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    folder = Path(tempfile.mkdtemp(prefix="vl_export_"))
    txt_path = folder / f"vl_inventory_{stamp}.txt"
    lines = [
        f"Victory Live inventory export {stamp}",
        f"Listed: {len(listed)}    Unlisted: {len(unlisted)}",
        "",
        "=== LISTED ===",
        "\t".join(_EXPORT_HEADERS),
    ]
    lines.extend("\t".join(str(v) for v in _export_values(r)) for r in listed)
    lines.extend(["", "=== UNLISTED ===", "\t".join(_EXPORT_HEADERS)])
    lines.extend("\t".join(str(v) for v in _export_values(r)) for r in unlisted)
    txt_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return txt_path, len(listed), len(unlisted)


def _write_inventory_xlsx(rows: list, folder: Path, stamp_name: str) -> Path:
    listed, unlisted = _split_listed(rows)
    from openpyxl import Workbook

    xlsx_path = folder / stamp_name
    wb = Workbook()
    listed_sheet = wb.active
    listed_sheet.title = "Listed"
    unlisted_sheet = wb.create_sheet("Unlisted")
    for sheet, sheet_rows in ((listed_sheet, listed), (unlisted_sheet, unlisted)):
        sheet.append(_EXPORT_HEADERS)
        for row in sheet_rows:
            sheet.append(_export_values(row))
        sheet.auto_filter.ref = sheet.dimensions
        sheet.freeze_panes = "A2"
    wb.save(xlsx_path)
    return xlsx_path


async def listed_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not await _deny_unless_admin(update):
        return
    msg = await update.message.reply_text("Building listed and unlisted export…")
    rows = db.export_active_groups()
    try:
        txt_path, n_listed, n_unlisted = await asyncio.to_thread(_write_inventory_txt, rows)
    except Exception as exc:
        await msg.edit_text(f"Export failed: {exc}")
        return
    caption = f"Listed {n_listed} · Unlisted {n_unlisted}"
    xlsx_name = txt_path.with_suffix(".xlsx").name
    try:
        xlsx_path = await asyncio.to_thread(_write_inventory_xlsx, rows, txt_path.parent, xlsx_name)
    except ImportError:
        with txt_path.open("rb") as txt_file:
            await update.message.reply_document(txt_file, filename=txt_path.name, caption=caption)
        await msg.edit_text(caption + "\nExcel package is missing on the server.")
        return
    except Exception as exc:
        with txt_path.open("rb") as txt_file:
            await update.message.reply_document(txt_file, filename=txt_path.name, caption=caption)
        await msg.edit_text(f"{caption}\nExcel failed: {exc}")
        return
    with txt_path.open("rb") as txt_file, xlsx_path.open("rb") as xlsx_file:
        await update.message.reply_document(txt_file, filename=txt_path.name, caption=caption)
        await update.message.reply_document(xlsx_file, filename=xlsx_path.name, caption=caption)
    await msg.edit_text(caption)


def _norm_date(token: str) -> str | None:
    token = token.strip()
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%Y/%m/%d"):
        try:
            return datetime.strptime(token, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


def _parse_date_range(args: list[str]) -> tuple[str, str, str] | None:
    """Return (lo, hi, label) as inclusive YYYY-MM-DD strings, or None if unparseable."""
    tokens = [a for a in args if a.strip()]
    if not tokens:
        return None
    if len(tokens) == 1:
        t = tokens[0].strip()
        if re.fullmatch(r"\d{4}", t):  # whole year
            return f"{t}-01-01", f"{t}-12-31", t
        if re.fullmatch(r"\d{4}-\d{2}", t):  # whole month
            year, month = int(t[:4]), int(t[5:7])
            last = calendar.monthrange(year, month)[1]
            return f"{t}-01", f"{t}-{last:02d}", t
        one = _norm_date(t)
        if one:
            return one, one, one
        return None
    lo = _norm_date(tokens[0])
    hi = _norm_date(tokens[1])
    if not lo or not hi:
        return None
    if lo > hi:
        lo, hi = hi, lo
    return lo, hi, f"{lo} → {hi}"


def _sort_by_date(rows: list) -> list:
    return sorted(
        rows,
        key=lambda r: (
            (r["production_date"] or "")[:10],
            r["production_name"] or "",
            r["venue"] or "",
            r["section"] or "",
            r["row_name"] or "",
        ),
    )


def _vl_group_in_range(row, lo: str, hi: str) -> bool:
    pd = (row["production_date"] or "")[:10]
    if lo <= pd <= hi:
        return True
    iso = event_date_to_iso(row["production_date"])
    return bool(iso and lo <= iso <= hi)


def _filter_stock_by_range(rows: list, lo: str, hi: str) -> list:
    out: list = []
    for row in rows:
        iso = event_date_to_iso(row["event_date"])
        if iso and lo <= iso <= hi:
            out.append(row)
    return out


_JUNK_BARCODE = {"", "no barc", "nobarc", "none", "n/a", "missing", "not found"}
_STOCKFROM_PRESET_LABELS = {
    "all": "ALL",
    "barcode": "Barcode only",
    "barcode_link": "Barcode + tix link",
    "bc_tix": "Barcode / tix link",
    "barcode_normal": "Barcode + normal",
    "app": "App events only",
    "no_info": "No info only (non-app)",
}


def _parse_one_event_date(text: str) -> str | None:
    text = (text or "").strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    for fmt in ("%b %d %Y", "%B %d %Y", "%b %d, %Y", "%B %d, %Y"):
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    for fmt in ("%b %d", "%B %d"):
        try:
            parsed = datetime.strptime(text, fmt)
            return parsed.replace(year=2026).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return _norm_date(text)


def _stockfrom_range_label(lo: str, hi: str) -> str:
    return f"{lo} → {hi}"


def _stockfrom_non_sf_in_range(lo: str, hi: str) -> tuple[list, dict]:
    """Non-SF tickets in [lo, hi] (SQL-filtered; no full-table load)."""
    on_sf_ids, sf_run_id = db.stock_ticket_ids_on_stagefront()
    in_range = db.export_stock_stockfrom_range(lo, hi)
    meta = {
        "sf_split_run_id": sf_run_id,
        "on_sf_ids": len(on_sf_ids),
        "off_sf_total": 0,
        "in_range_before_preset": len(in_range),
    }
    return in_range, meta


def _stockfrom_preset_counts(lo: str, hi: str) -> dict[str, int]:
    """Non-SF ticket counts in range per export preset (for filter button labels)."""
    in_range, meta = _stockfrom_non_sf_in_range(lo, hi)
    out: dict[str, int] = {
        preset: len(_filter_stock_preset(in_range, preset))
        for preset in (
            "all",
            "barcode",
            "barcode_link",
            "bc_tix",
            "barcode_normal",
            "app",
            "no_info",
        )
    }
    out["in_range"] = int(meta["in_range_before_preset"])
    return out


def _stockfrom_scan_summary(lo: str, hi: str) -> tuple[dict[str, int], dict[str, int], dict]:
    return db.stockfrom_scan_summary(lo, hi)


def _stockfrom_keyboard(counts: dict[str, int] | None = None) -> InlineKeyboardMarkup:
    c = counts or {}
    def _n(preset: str) -> str:
        n = c.get(preset)
        return f" ({n})" if n is not None else ""

    return InlineKeyboardMarkup(
        [
            [InlineKeyboardButton(f"ALL{_n('all')}", callback_data="stf:all")],
            [InlineKeyboardButton(f"Barcode only{_n('barcode')}", callback_data="stf:bc")],
            [InlineKeyboardButton(f"Barcode + tix link{_n('barcode_link')}", callback_data="stf:bcl")],
            [InlineKeyboardButton(f"Barcode / tix link{_n('bc_tix')}", callback_data="stf:bxt")],
            [InlineKeyboardButton(f"Barcode + normal{_n('barcode_normal')}", callback_data="stf:bcn")],
            [InlineKeyboardButton(f"App events only{_n('app')}", callback_data="stf:app")],
            [InlineKeyboardButton(f"No info only (non-app){_n('no_info')}", callback_data="stf:noinfo")],
        ]
    )


async def _stockfrom_offer_filters(update: Update, context: ContextTypes.DEFAULT_TYPE, lo: str, hi: str) -> None:
    label = _stockfrom_range_label(lo, hi)
    context.user_data["stockfrom_pending"] = {"lo": lo, "hi": hi, "label": label}
    context.user_data.pop("stockfrom_lo", None)
    status = await update.effective_message.reply_text(
        f"Event dates: <b>{html.escape(label)}</b>\nScanning ticket database…",
        parse_mode=ParseMode.HTML,
    )
    try:
        counts, kinds, meta = await asyncio.to_thread(_stockfrom_scan_summary, lo, hi)
    except Exception as exc:
        log.exception("stockfrom scan failed lo=%s hi=%s", lo, hi)
        try:
            await status.edit_text(f"Scan failed: {html.escape(str(exc))}")
        except Exception:
            await update.effective_message.reply_text(f"Scan failed: {exc}")
        return
    sf_run = meta.get("sf_split_run_id")
    on_sf = int(meta.get("on_sf_ids") or 0)
    sf_note = (
        f"SF split run #{sf_run}: excluding {on_sf} on-Stagefront rows."
        if sf_run
        else "No SF split run yet — export is all bot stock in range (not SF-filtered)."
    )
    kind_line = ""
    seats = int(counts.get("seats") or counts["all"])
    db_total = int(meta.get("db_total") or 0)
    if kinds:
        kind_line = (
            "\nIn range: "
            + ", ".join(f"{k} {n}" for k, n in sorted(kinds.items()))
            + f" · <b>{counts['all']}</b> rows · <b>{seats}</b> seats (qty)"
        )
    else:
        kind_line = f"\n<b>{counts['all']}</b> rows · <b>{seats}</b> seats (qty)"
    if db_total and counts["all"] < db_total // 3:
        kind_line += (
            f"\n<i>DB has {db_total:,} tickets total; only rows with a parseable "
            f"event date in this range count. Many imports store venue/artist in "
            f"the date field — re-import or widen dates if exports look low.</i>"
        )
    body = (
        f"Event dates: <b>{html.escape(label)}</b>\n{html.escape(sf_note)}{kind_line}\n\nChoose export:"
    )
    markup = _stockfrom_keyboard(counts)
    try:
        await status.edit_text(body, reply_markup=markup, parse_mode=ParseMode.HTML)
    except Exception:
        log.exception("stockfrom edit_text failed — sending new message")
        await update.effective_message.reply_text(
            body, reply_markup=markup, parse_mode=ParseMode.HTML
        )


def _stock_has_barcode(row) -> bool:
    bc = (row["barcode"] or "").strip().lower()
    return bool(bc) and bc not in _JUNK_BARCODE


def _stock_has_tix_link(row) -> bool:
    url = (row["ticket_url"] or "").strip() or (row["transfer_url"] or "").strip()
    if url:
        return True
    tok = (row["secure_token"] or "").strip()
    return tok.startswith("http") or "tixx." in tok.lower()


def _stock_has_barcode_or_tix(row) -> bool:
    return _stock_has_barcode(row) or _stock_has_tix_link(row)


def _stock_is_no_info(row) -> bool:
    """Checker EVENT w/ TICKETS (none) — excludes app-style rows (use App events only)."""
    if (row["kind"] or "").strip().lower() == "app":
        return False
    if _stock_has_barcode(row) or _stock_has_tix_link(row):
        return False
    if (row["section"] or "").strip():
        return False
    if (row["row_name"] or "").strip():
        return False
    if (row["seat"] or "").strip():
        return False
    return bool((row["event_name"] or "").strip() or (row["event_date"] or "").strip())


def _filter_stock_preset(rows: list, preset: str) -> list:
    if preset == "barcode":
        return [r for r in rows if _stock_has_barcode(r)]
    if preset == "barcode_link":
        return [r for r in rows if _stock_has_barcode(r) and _stock_has_tix_link(r)]
    if preset == "bc_tix":
        return [r for r in rows if _stock_has_barcode_or_tix(r)]
    if preset == "barcode_normal":
        return [r for r in rows if (r["kind"] or "") in ("barcode", "normal")]
    if preset == "app":
        return [r for r in rows if (r["kind"] or "") == "app"]
    if preset == "no_info":
        return [r for r in rows if _stock_is_no_info(r)]
    return list(rows)


def _build_non_sf_stock_export(lo: str, hi: str, preset: str) -> tuple[list, dict]:
    in_range, meta = _stockfrom_non_sf_in_range(lo, hi)
    filtered = _sort_stock_rows(_filter_stock_preset(in_range, preset))
    meta = {
        **meta,
        "exported": len(filtered),
        "preset": preset,
    }
    return filtered, meta


def _sort_stock_rows(rows: list) -> list:
    kind_order = {"app": 0, "barcode": 1, "transfer": 2, "normal": 3}
    return sorted(
        rows,
        key=lambda r: (
            kind_order.get(r["kind"] or "", 9),
            event_date_to_iso(r["event_date"]) or (r["event_date"] or ""),
            r["event_name"] or "",
            r["email"] or "",
            r["section"] or "",
            r["row_name"] or "",
            r["seat"] or "",
        ),
    )


def _stock_kind_counts(rows: list) -> dict[str, int]:
    counts: dict[str, int] = {}
    for row in rows:
        k = row["kind"] or "other"
        counts[k] = counts.get(k, 0) + 1
    return counts


def _write_bydate_txt(vl_rows: list, stock_rows: list, label: str, folder: Path, stamp: str) -> Path:
    txt_path = folder / f"vl_bydate_{stamp}.txt"
    listed = sum(1 for r in vl_rows if r["is_shared"])
    stock_counts = _stock_kind_counts(stock_rows)
    lines = [
        f"Victory Live + ticket database · {label}",
        f"VL groups: {len(vl_rows)} (listed {listed}, unlisted {len(vl_rows) - listed})",
        f"Stored tickets: {len(stock_rows)}"
        + (
            " (" + ", ".join(f"{k} {n}" for k, n in sorted(stock_counts.items())) + ")"
            if stock_counts
            else ""
        ),
        "",
    ]
    if vl_rows:
        lines.extend(["=== VICTORY LIVE INVENTORY ===", "\t".join(_EXPORT_HEADERS)])
        lines.extend("\t".join(str(v) for v in _export_values(r)) for r in vl_rows)
        lines.append("")
    if stock_rows:
        by_kind: dict[str, list] = {}
        for row in stock_rows:
            by_kind.setdefault(row["kind"] or "other", []).append(row)
        for kind in ("app", "barcode", "transfer", "normal"):
            chunk = by_kind.get(kind) or []
            if not chunk:
                continue
            lines.extend([f"=== STORED · {kind.upper()} ({len(chunk)}) ===", "\t".join(_STOCK_HEADERS)])
            lines.extend("\t".join(str(v).replace("\t", " ") for v in _stock_values(r)) for r in chunk)
            lines.append("")
    txt_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return txt_path


def _append_stock_sheet(wb, title: str, rows: list) -> None:
    sheet = wb.create_sheet(title[:31])
    sheet.append(_STOCK_HEADERS)
    for row in rows:
        sheet.append(_stock_values(row))
    if sheet.max_row > 1:
        sheet.auto_filter.ref = sheet.dimensions
        sheet.freeze_panes = "A2"


def _write_bydate_xlsx(vl_rows: list, stock_rows: list, folder: Path, name: str) -> Path:
    from openpyxl import Workbook

    xlsx_path = folder / name
    wb = Workbook()
    first = True
    if vl_rows:
        sheet = wb.active
        sheet.title = "VL Inventory"
        sheet.append(_EXPORT_HEADERS)
        for row in vl_rows:
            sheet.append(_export_values(row))
        sheet.auto_filter.ref = sheet.dimensions
        sheet.freeze_panes = "A2"
        first = False
    by_kind: dict[str, list] = {}
    for row in stock_rows:
        by_kind.setdefault(row["kind"] or "other", []).append(row)
    if stock_rows and first:
        wb.active.title = "All tickets"
        wb.active.append(_STOCK_HEADERS)
        for row in stock_rows:
            wb.active.append(_stock_values(row))
        wb.active.auto_filter.ref = wb.active.dimensions
        wb.active.freeze_panes = "A2"
        first = False
    elif stock_rows:
        _append_stock_sheet(wb, "All tickets", stock_rows)
    for kind, label in (
        ("app", "App events"),
        ("barcode", "Barcodes"),
        ("transfer", "Transfers"),
        ("normal", "Seats"),
    ):
        chunk = by_kind.get(kind) or []
        if chunk:
            _append_stock_sheet(wb, label, chunk)
    if first:
        sheet = wb.active
        sheet.title = "Empty"
        sheet.append(["No rows for this date range"])
    wb.save(xlsx_path)
    return xlsx_path


async def bydate_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not await _deny_unless_admin(update):
        return
    rng = _parse_date_range(context.args or [])
    if not rng:
        await update.message.reply_text(
            "Usage: /bydate <date> or /bydate <from> <to>\n"
            "Examples:\n"
            "  /bydate 2026-10-04\n"
            "  /bydate 2026-10-04 2026-10-20\n"
            "  /bydate 2026-10   (whole month)\n"
            "  /bydate 2026      (whole year)\n"
            "Accepts YYYY-MM-DD or MM/DD/YYYY.\n"
            "Includes Victory Live active groups + stored tickets (app events, barcodes, seats)."
        )
        return
    lo, hi, label = rng
    msg = await update.message.reply_text(f"Filtering for {label}…")
    vl_rows, stock_all = await asyncio.gather(
        asyncio.to_thread(db.export_active_groups),
        asyncio.to_thread(db.export_stock),
    )
    vl_rows = _sort_by_date([r for r in vl_rows if _vl_group_in_range(r, lo, hi)])
    stock_rows = _sort_stock_rows(_filter_stock_by_range(stock_all, lo, hi))
    if not vl_rows and not stock_rows:
        await msg.edit_text(
            f"Nothing for {label} in Victory Live inventory or the ticket database."
        )
        return
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    folder = Path(tempfile.mkdtemp(prefix="vl_bydate_"))
    sc = _stock_kind_counts(stock_rows)
    stock_bit = f" · stock {len(stock_rows)}"
    if sc:
        stock_bit += " (" + ", ".join(f"{k} {n}" for k, n in sorted(sc.items())) + ")"
    caption = f"{label} · VL {len(vl_rows)} groups{stock_bit}"
    try:
        txt_path = await asyncio.to_thread(_write_bydate_txt, vl_rows, stock_rows, label, folder, stamp)
    except Exception as exc:
        await msg.edit_text(f"Export failed: {exc}")
        return
    try:
        xlsx_path = await asyncio.to_thread(
            _write_bydate_xlsx, vl_rows, stock_rows, folder, f"vl_bydate_{stamp}.xlsx"
        )
    except ImportError:
        with txt_path.open("rb") as txt_file:
            await update.message.reply_document(txt_file, filename=txt_path.name, caption=caption)
        await msg.edit_text(caption + "\nExcel package is missing on the server.")
        return
    except Exception as exc:
        with txt_path.open("rb") as txt_file:
            await update.message.reply_document(txt_file, filename=txt_path.name, caption=caption)
        await msg.edit_text(f"{caption}\nExcel failed: {exc}")
        return
    with txt_path.open("rb") as txt_file, xlsx_path.open("rb") as xlsx_file:
        await update.message.reply_document(txt_file, filename=txt_path.name, caption=caption)
        await update.message.reply_document(xlsx_file, filename=xlsx_path.name, caption=caption)
    await msg.edit_text(caption)


async def stats_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not await _deny_unless_admin(update):
        return
    msg = await update.message.reply_text(
        _db_stats_text(db.summary()) + "\n\nCounting Victory Live…"
    )
    try:
        live = await asyncio.to_thread(
            live_inventory_stats,
            days_back=SYNC_DAYS_BACK,
            days_forward=SYNC_DAYS_FORWARD,
            affiliate_group_id=AFFILIATE_GROUP_ID,
            use_proxy=USE_PROXY,
        )
    except Exception as exc:
        await msg.edit_text(_db_stats_text(db.summary()) + f"\n\nVictory Live count failed: {exc}")
        return
    if not live.get("ok"):
        await msg.edit_text(
            _db_stats_text(db.summary()) + f"\n\nVictory Live count failed: {live.get('error')}"
        )
        return
    text = (
        _db_stats_text(db.summary())
        + "\n\nVictory Live (live)\n"
        + f"Window: {live['from']} → {live['to']}\n"
        + f"Productions: {live['productions']}\n"
        + f"Active groups: {live['active']} · tickets {live['qty']}\n"
        + f"Listed: {live['listed']} · tickets {live['listed_qty']} · list {_money(live['listed_list_value'])}\n"
        + f"Unlisted: {live['unlisted']}\n"
        + f"Active list value: {_money(live['list_value'])}\n"
        + f"Active cost value: {_money(live['cost_value'])}"
    )
    await msg.edit_text(text)


async def fetchibrahim_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not await _deny_unless_admin(update):
        return
    if not update.message:
        return
    msg = await update.message.reply_text(
        "Fetching Ibrahim lister stats (bot DB + Victory Live)…"
    )
    try:
        report = await asyncio.to_thread(
            build_ibrahim_report,
            use_proxy=USE_PROXY,
            affiliate_group_id=AFFILIATE_GROUP_ID,
        )
    except Exception as exc:
        log.exception("fetchibrahim failed")
        await msg.edit_text(f"Fetch failed: {html.escape(str(exc))}", parse_mode=ParseMode.HTML)
        return
    await msg.edit_text(format_ibrahim_report_html(report), parse_mode=ParseMode.HTML)


def _stock_summary_text() -> str:
    s = db.stock_summary()
    lines = [
        "Ticket database",
        f"Accounts: {s['accounts']}",
        f"Ticket rows: {s['tickets']}",
    ]
    labels = {
        "app": "App events",
        "barcode": "Barcodes",
        "transfer": "Transfers",
        "normal": "Normal seats",
    }
    for kind, label in labels.items():
        info = s["by_kind"].get(kind)
        if not info:
            lines.append(f"{label}: 0")
            continue
        extra = f" · with barcode {info['barcodes']}" if kind != "app" else ""
        lines.append(f"{label}: {info['rows']}{extra}")
    return "\n".join(lines)


async def stock_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not await _deny_unless_admin(update):
        return
    await update.message.reply_text(
        _stock_summary_text()
        + "\n\nSend a .txt or .csv to import.\n"
        "App events, barcode lines, transfer links, and seat blocks are stored in full."
    )


def _stock_card(row) -> str:
    where = " ".join(
        part
        for part in (
            f"Sec {row['section']}" if row["section"] else "",
            f"Row {row['row_name']}" if row["row_name"] else "",
            f"Seat {row['seat']}" if row["seat"] else "",
        )
        if part
    )
    bits = [f"{row['kind']} · {row['event_name'] or '—'}"]
    bits.append(f"{row['event_date'] or '—'} · {row['venue'] or '—'}")
    if row["app_name"]:
        bits.append(f"APP: {row['app_name']} · qty {row['quantity'] or 0}")
    if where:
        bits.append(where)
    if row["barcode"]:
        bits.append(f"Barcode: {row['barcode']}")
    if row["transfer_status"]:
        bits.append(row["transfer_status"])
    if row["transfer_url"]:
        bits.append(row["transfer_url"])
    bits.append(row["email"])
    return "\n".join(bits)


async def stockfind_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not await _deny_unless_admin(update):
        return
    q = " ".join(context.args or []).strip()
    if not q:
        await update.message.reply_text("Usage: /stockfind bruno or /stockfind CLUB25 or /stockfind email")
        return
    rows = db.search_stock(q, limit=15)
    if not rows:
        await update.message.reply_text("No stored tickets match.")
        return
    text = "\n\n".join(_stock_card(r) for r in rows)
    if len(text) > 4000:
        text = text[:4000] + "\n…"
    await update.message.reply_text(text)


_STOCK_HEADERS = [
    "kind",
    "email",
    "password",
    "phone",
    "full_name",
    "event",
    "date",
    "venue",
    "app",
    "section",
    "row",
    "seat",
    "qty",
    "barcode",
    "secure_token",
    "transfer_status",
    "transfer_url",
    "order_id",
    "price",
    "ticket_url",
    "payment_methods",
    "source",
]


def _stock_values(row) -> list:
    return [
        row["kind"] or "",
        row["email"] or "",
        row["password"] or "",
        row["phone"] or "",
        row["full_name"] or "",
        row["event_name"] or "",
        row["event_date"] or "",
        row["venue"] or "",
        row["app_name"] or "",
        row["section"] or "",
        row["row_name"] or "",
        row["seat"] or "",
        row["quantity"] if row["quantity"] is not None else "",
        row["barcode"] or "",
        row["secure_token"] or "",
        row["transfer_status"] or "",
        row["transfer_url"] or "",
        row["order_id"] or "",
        row["price"] or "",
        row["ticket_url"] or "",
        row["payment_methods"] or "",
        row["source_file"] or "",
    ]


def _write_stock_export(rows: list) -> tuple[Path, Path]:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    folder = Path(tempfile.mkdtemp(prefix="vl_stock_"))
    txt_path = folder / f"ticket_stock_{stamp}.txt"
    lines = ["\t".join(_STOCK_HEADERS)]
    lines.extend("\t".join(str(v).replace("\t", " ") for v in _stock_values(r)) for r in rows)
    txt_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    from openpyxl import Workbook

    xlsx_path = folder / f"ticket_stock_{stamp}.xlsx"
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Tickets"
    sheet.append(_STOCK_HEADERS)
    for row in rows:
        sheet.append(_stock_values(row))
    sheet.freeze_panes = "A2"
    wb.save(xlsx_path)
    return txt_path, xlsx_path


def _write_stock_noinfo_simple_txt(rows: list, folder: Path, stamp: str) -> tuple[Path, int]:
    """email:password + event date + event name (no seat/barcode rows)."""
    lines: list[str] = ["email:password\tevent_date\tevent_name\tvenue\tapp_name"]
    for row in rows:
        email = (row["email"] or "").strip()
        password = (row["password"] or "").strip()
        if not email or not password or "@" not in email:
            continue
        dt = event_date_to_iso(row["event_date"]) or (row["event_date"] or "").strip()
        ev = (row["event_name"] or "").replace("\t", " ").strip()
        venue = (row["venue"] or "").replace("\t", " ").strip()
        app = (row["app_name"] or "").replace("\t", " ").strip()
        lines.append(f"{email}:{password}\t{dt}\t{ev}\t{venue}\t{app}")
    path = folder / f"ticket_stock_{stamp}_noinfo.txt"
    path.write_text("\n".join(lines) + ("\n" if len(lines) > 1 else ""), encoding="utf-8")
    return path, max(0, len(lines) - 1)


def _write_stock_creds_txt(rows: list, folder: Path, stamp: str) -> tuple[Path, int]:
    """One email:password per line (unique accounts from export rows)."""
    seen: set[str] = set()
    lines: list[str] = []
    for row in rows:
        email = (row["email"] or "").strip().lower()
        password = (row["password"] or "").strip()
        if not email or not password or "@" not in email:
            continue
        if email in seen:
            continue
        seen.add(email)
        lines.append(f"{email}:{password}")
    creds_path = folder / f"ticket_stock_{stamp}_creds.txt"
    creds_path.write_text("\n".join(lines) + ("\n" if lines else ""), encoding="utf-8")
    return creds_path, len(lines)


_STOCKFROM_PRESET_MAP = {
    "stf:all": "all",
    "stf:bc": "barcode",
    "stf:bcl": "barcode_link",
    "stf:bxt": "bc_tix",
    "stf:bcn": "barcode_normal",
    "stf:app": "app",
    "stf:noinfo": "no_info",
}

_STOCKFROM_LOCKS: dict[int, asyncio.Lock] = {}


def _stockfrom_user_lock(user_id: int) -> asyncio.Lock:
    lock = _STOCKFROM_LOCKS.get(user_id)
    if lock is None:
        lock = asyncio.Lock()
        _STOCKFROM_LOCKS[user_id] = lock
    return lock


def _stockfrom_clear_await(context: ContextTypes.DEFAULT_TYPE, telegram_id: int | None = None) -> None:
    context.user_data.pop("stockfrom_await", None)
    context.user_data.pop("stockfrom_lo", None)
    if telegram_id is not None:
        db.clear_stockfrom_session(telegram_id)


def _stockfrom_persist(context: ContextTypes.DEFAULT_TYPE, telegram_id: int) -> None:
    step = context.user_data.get("stockfrom_await")
    if not step:
        db.clear_stockfrom_session(telegram_id)
        return
    lo = context.user_data.get("stockfrom_lo")
    db.save_stockfrom_session(telegram_id, str(step), str(lo) if lo else None)


def _stockfrom_restore(context: ContextTypes.DEFAULT_TYPE, telegram_id: int) -> str | None:
    """Merge DB wizard state into user_data (survives bot restart / lost user_data)."""
    if context.user_data.get("stockfrom_await"):
        return str(context.user_data["stockfrom_await"])
    row = db.get_stockfrom_session(telegram_id)
    if not row:
        return None
    step = str(row["step"] or "")
    if not step:
        return None
    context.user_data["stockfrom_await"] = step
    if row["lo"]:
        context.user_data["stockfrom_lo"] = str(row["lo"])
    return step


def _clear_vl_listing_user_state(context: ContextTypes.DEFAULT_TYPE) -> None:
    context.user_data.pop("vl_list", None)
    context.user_data.pop("vl_list_await", None)
    for key in list(context.user_data.keys()):
        if isinstance(key, tuple) and key and key[0] == "vl_listing":
            context.user_data.pop(key, None)


async def stockfrom_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not update.effective_user or not update.message:
        return
    if not _admin(update.effective_user.id):
        await update.message.reply_text("Admin only.")
        return
    _clear_vl_listing_user_state(context)
    uid = update.effective_user.id
    _stockfrom_clear_await(context, uid)
    args = [a for a in (context.args or []) if str(a).strip()]
    if len(args) == 1:
        lo = _parse_one_event_date(args[0])
        if lo:
            await _stockfrom_offer_filters(update, context, lo, lo)
            return
        await update.message.reply_text(
            "Could not parse that date.\n"
            "Example: <code>/stockfrom 2026-10-02</code> (one day)\n"
            "Or: <code>/stockfrom 2026-11-23 2027-12-31</code>",
            parse_mode=ParseMode.HTML,
        )
        return
    if len(args) >= 2:
        lo = _parse_one_event_date(args[0])
        hi = _parse_one_event_date(" ".join(args[1:]) if len(args) > 2 else args[1])
        if not lo or not hi:
            await update.message.reply_text(
                "Could not parse both dates.\n"
                "Example: /stockfrom 2026-11-23 2027-12-31"
            )
            return
        if lo > hi:
            lo, hi = hi, lo
        await _stockfrom_offer_filters(update, context, lo, hi)
        return
    context.user_data["stockfrom_await"] = "start"
    _stockfrom_persist(context, uid)
    await update.message.reply_text(
        "Non-Stagefront export — step 1/2\n\n"
        "Send the <b>start</b> event date (inclusive), e.g. <code>2026-11-23</code>\n"
        "Or one day in one shot: <code>/stockfrom 2026-10-02</code>",
        parse_mode=ParseMode.HTML,
    )


def _stockfrom_parse_range(text: str) -> tuple[str, str] | None:
    """Parse one or two event dates from user text."""
    raw = (text or "").strip()
    if not raw:
        return None
    if raw.lower() in ("same", ".", "="):
        return None
    parts = raw.split()
    if len(parts) >= 2:
        lo = _parse_one_event_date(parts[0])
        hi = _parse_one_event_date(" ".join(parts[1:]))
        if lo and hi:
            if lo > hi:
                lo, hi = hi, lo
            return lo, hi
    one = _parse_one_event_date(raw)
    if one:
        return one, one
    return None


async def stockfrom_date_reply(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    """Date steps for /stockfrom (user_data, not ConversationHandler — safe with concurrent_updates)."""
    msg = update.effective_message
    if not msg or not msg.text or not update.effective_user:
        return
    if not _admin(update.effective_user.id):
        if context.user_data.get("stockfrom_await"):
            _stockfrom_clear_await(context, update.effective_user.id)
        return

    uid = update.effective_user.id
    step = _stockfrom_restore(context, uid)
    if not step:
        return

    scan_range: tuple[str, str] | None = None
    async with _stockfrom_user_lock(uid):
        step = _stockfrom_restore(context, uid)
        if not step:
            return
        text = msg.text.strip()
        log.info("stockfrom date step=%s user=%s text=%s", step, uid, text[:40])
        if step == "start":
            parts = text.split()
            if len(parts) >= 2:
                pair = _stockfrom_parse_range(text)
                if not pair:
                    await msg.reply_text(
                        "Invalid range. Example: <code>2026-11-23 2027-06-30</code>",
                        parse_mode=ParseMode.HTML,
                    )
                    raise ApplicationHandlerStop()
                context.user_data.pop("stockfrom_await", None)
                context.user_data.pop("stockfrom_lo", None)
                db.clear_stockfrom_session(uid)
                scan_range = pair
            else:
                lo = _parse_one_event_date(text)
                if not lo:
                    await msg.reply_text(
                        "Invalid date. Example: <code>2026-11-23</code>\n"
                        "One day: <code>/stockfrom 2026-10-02</code>",
                        parse_mode=ParseMode.HTML,
                    )
                    raise ApplicationHandlerStop()
                context.user_data["stockfrom_lo"] = lo
                context.user_data["stockfrom_await"] = "end"
                _stockfrom_persist(context, uid)
                await msg.reply_text(
                    f"Start: <b>{html.escape(lo)}</b> — step 2/2\n\n"
                    "Send the <b>end</b> date (inclusive), or <code>same</code> for one day:",
                    parse_mode=ParseMode.HTML,
                )
                raise ApplicationHandlerStop()
        elif step == "end":
            lo = context.user_data.get("stockfrom_lo")
            if not lo:
                row = db.get_stockfrom_session(uid)
                lo = str(row["lo"]) if row and row["lo"] else None
                if lo:
                    context.user_data["stockfrom_lo"] = lo
            if not lo:
                _stockfrom_clear_await(context, uid)
                await msg.reply_text("Session lost — send /stockfrom again.")
                raise ApplicationHandlerStop()
            parsed = _parse_one_event_date(text)
            if text.lower() in ("same", ".", "="):
                hi = lo
            elif not parsed:
                context.user_data["stockfrom_await"] = "end"
                _stockfrom_persist(context, uid)
                await msg.reply_text(
                    "Invalid end date. Example: 2027-06-30 or <code>same</code>",
                    parse_mode=ParseMode.HTML,
                )
                raise ApplicationHandlerStop()
            else:
                hi = parsed
            if lo > hi:
                lo, hi = hi, lo
            context.user_data.pop("stockfrom_await", None)
            context.user_data.pop("stockfrom_lo", None)
            db.clear_stockfrom_session(uid)
            scan_range = (lo, hi)

    if scan_range:
        await _stockfrom_offer_filters(update, context, scan_range[0], scan_range[1])
        raise ApplicationHandlerStop()


async def stockfrom_button(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    q = update.callback_query
    if not q or not q.data or not q.from_user:
        return
    if not q.data.startswith("stf:"):
        return
    if not _admin(q.from_user.id):
        await q.answer("Admin only", show_alert=True)
        return
    preset = _STOCKFROM_PRESET_MAP.get(q.data)
    if not preset:
        await q.answer()
        return
    pending = context.user_data.get("stockfrom_pending")
    if not pending:
        await q.answer("Expired — send /stockfrom again", show_alert=True)
        return
    await q.answer()
    lo, hi, label = pending["lo"], pending["hi"], pending["label"]
    preset_label = _STOCKFROM_PRESET_LABELS.get(preset, preset)
    try:
        await q.edit_message_text(f"Building {preset_label} · {label}…")
    except Exception:
        pass
    rows, meta = await asyncio.to_thread(_build_non_sf_stock_export, lo, hi, preset)
    if not rows:
        counts = await asyncio.to_thread(_stockfrom_preset_counts, lo, hi)
        hint = ""
        if preset != "all" and counts.get("all"):
            hint = f"\n\nTry <b>ALL ({counts['all']})</b> — other kinds may be in this range."
        await context.bot.send_message(
            chat_id=q.message.chat_id,
            text=f"No rows for {preset_label} · {label} (non-SF).{hint}",
            parse_mode=ParseMode.HTML,
        )
        return
    txt_path, xlsx_path = await asyncio.to_thread(_write_stock_export, rows)
    stamp = txt_path.stem.replace("ticket_stock_", "", 1)
    creds_path, creds_count = await asyncio.to_thread(
        _write_stock_creds_txt, rows, txt_path.parent, stamp
    )
    noinfo_path, noinfo_count = (None, 0)
    if preset == "no_info":
        noinfo_path, noinfo_count = await asyncio.to_thread(
            _write_stock_noinfo_simple_txt, rows, txt_path.parent, stamp
        )
    sf_bit = f" · SF run #{meta['sf_split_run_id']}" if meta.get("sf_split_run_id") else ""
    caption = (
        f"{preset_label} · {label}\n"
        f"Non-SF only · {len(rows)} tickets{sf_bit}"
    )
    with txt_path.open("rb") as txt_file, xlsx_path.open("rb") as xlsx_file:
        await context.bot.send_document(
            chat_id=q.message.chat_id,
            document=txt_file,
            filename=txt_path.name,
            caption=caption,
        )
        await context.bot.send_document(
            chat_id=q.message.chat_id,
            document=xlsx_file,
            filename=xlsx_path.name,
            caption=caption,
        )
    if noinfo_path and noinfo_path.is_file() and noinfo_count:
        noinfo_caption = (
            f"{preset_label} · {label}\n"
            f"email:pw + event date · {noinfo_count} rows"
        )
        with noinfo_path.open("rb") as noinfo_file:
            await context.bot.send_document(
                chat_id=q.message.chat_id,
                document=noinfo_file,
                filename=noinfo_path.name,
                caption=noinfo_caption,
            )
    if creds_path and creds_path.is_file() and creds_count and preset != "no_info":
        creds_caption = f"{preset_label} · {label}\nemail:pw · {creds_count} accounts"
        with creds_path.open("rb") as creds_file:
            await context.bot.send_document(
                chat_id=q.message.chat_id,
                document=creds_file,
                filename=creds_path.name,
                caption=creds_caption,
            )
    try:
        await q.edit_message_text(f"Sent {len(rows)} tickets · {preset_label} · {label}")
    except Exception:
        pass
    context.user_data.pop("stockfrom_pending", None)


async def stockfile_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not await _deny_unless_admin(update):
        return
    rows = db.export_stock()
    if not rows:
        await update.message.reply_text("Ticket database is empty. Send a .txt or .csv to import.")
        return
    msg = await update.message.reply_text(f"Building export of {len(rows)} tickets…")
    txt_path, xlsx_path = await asyncio.to_thread(_write_stock_export, rows)
    caption = f"{len(rows)} stored tickets"
    with txt_path.open("rb") as txt_file, xlsx_path.open("rb") as xlsx_file:
        await update.message.reply_document(txt_file, filename=txt_path.name, caption=caption)
        await update.message.reply_document(xlsx_file, filename=xlsx_path.name, caption=caption)
    await msg.edit_text(caption)


def _import_stock_file(path: str, source: str) -> dict:
    file = Path(path)
    if file.suffix.lower() in (".xlsx", ".xlsm"):
        rows = vl_stock.parse_xlsx(str(file), source=source)
    else:
        rows = vl_stock.parse_inventory_text(
            file.read_text(encoding="utf-8-sig", errors="replace"),
            source=source,
        )
    if not rows:
        return {"parsed": 0, "inserted": 0, "updated": 0, "kinds": {}}
    stats = db.upsert_stock(rows)
    kinds: dict[str, int] = {}
    for row in rows:
        kinds[row["kind"]] = kinds.get(row["kind"], 0) + 1
    stats["parsed"] = len(rows)
    stats["kinds"] = kinds
    return stats


async def fulfill_file_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not await _deny_unless_admin(update):
        return
    st = await asyncio.to_thread(__import__("vl_fulfill_stock_db", fromlist=["stats"]).stats)
    context.user_data["fulfill_file_upload_pending"] = True
    await update.message.reply_text(
        "Send your <b>barcode inventory</b> file next (.xlsx or tab .txt).\n"
        "Format: checker export (kind, email, password, … event, date, venue, section, row, seat, barcode).\n"
        "Rows <b>merge</b> into the fulfill stock DB (same seat updates barcode; nothing is cleared).\n"
        f"Fulfill DB now: <b>{st['with_barcode']}</b> barcoded seats"
        + (f" · last update {st['last_updated']}" if st.get("last_updated") else ""),
        parse_mode=ParseMode.HTML,
    )


async def upload_batch_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    """TM account batch upload (/upload) — send .xlsx or .txt next."""
    if not await _deny_unless_admin(update):
        return
    if not _tm_loaded():
        await update.message.reply_text("TM worker module is not loaded on this server.")
        return
    context.user_data["tm_batch_upload_pending"] = True
    await update.message.reply_text(
        "Send a <b>.xlsx</b> or <b>.txt</b> <b>TM account batch</b> next "
        "(email/password + events for <code>/work</code>).\n"
        "Ticket stock txt/xlsx imports by sending the file without /upload.",
        parse_mode=ParseMode.HTML,
    )


def _tm_loaded() -> bool:
    try:
        from ticketmaster.tm_worker_bridge import tm_module_loaded

        return tm_module_loaded()
    except Exception:
        return False


async def dbstatus_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not await _deny_unless_admin(update):
        return
    lines = ["<b>Database status</b>", ""]
    st = db.stock_summary()
    lines.append("<b>VL ticket DB</b> (data/vl_inventory.db · stock_tickets)")
    lines.append(f"Accounts: {st['accounts']} · Tickets: {st['tickets']}")
    for kind, info in sorted((st.get("by_kind") or {}).items()):
        lines.append(f"  {kind}: {info['rows']} rows")
    on_sf, sf_run = db.stock_ticket_ids_on_stagefront()
    lines.append(f"SF split: run #{sf_run or '—'} · {len(on_sf)} tickets marked on-Stagefront")
    lines.append("")
    if _tm_loaded():
        from ticketmaster import bot as tm

        snap = tm.db.queue_snapshot()
        stats = tm.db.batch_stats()
        lines.append("<b>TM worker DB</b> (data/tm_worker/)")
        lines.append(
            f"Queue: pending {snap.get('pending_accounts', 0)} · "
            f"in progress {snap.get('in_progress_accounts', 0)} · "
            f"active workers {snap.get('active_workers', 0)}"
        )
        lines.append(f"Batches: {len(stats)}")
        for b in (stats or [])[:8]:
            lines.append(
                f"  #{b['id']} {html.escape(str(b['name'])[:40])} "
                f"done {b['done_count']}/{b['account_count']}"
            )
        if len(stats or []) > 8:
            lines.append(f"  … +{len(stats) - 8} more (/batches)")
    else:
        lines.append("<b>TM worker DB</b>: module not loaded")
    lines.append("")
    lines.append("Imports: TM batch → /upload or TM-shaped file · tickets → stock txt/xlsx")
    await update.message.reply_text("\n".join(lines), parse_mode=ParseMode.HTML)


async def stockdedupe_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not await _deny_unless_admin(update):
        return
    msg = await update.message.reply_text("Deduping stock_tickets…")
    result = await asyncio.to_thread(db.dedupe_stock)
    await msg.edit_text(
        f"Stock dedupe done.\n"
        f"Before: {result['before']} · Removed: {result['removed']} · Kept: {result['kept']}\n\n"
        + _stock_summary_text()
    )


async def admin_document_router(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    if not update.effective_user or not update.message or not update.message.document:
        return
    uid = update.effective_user.id
    if not _admin(uid):
        if _worker(uid):
            await update.message.reply_text(_WORKER_ONLY_HINT)
        return
    doc = update.message.document
    name = doc.file_name or "upload"
    dest_dir = Path(tempfile.mkdtemp(prefix="vl_upload_"))
    dest = dest_dir / name
    tg_file = await doc.get_file()
    await tg_file.download_to_drive(custom_path=str(dest))

    force_tm = context.user_data.pop("tm_batch_upload_pending", False)
    force_fulfill = context.user_data.pop("fulfill_file_upload_pending", False)
    if force_fulfill:
        await fulfill_stock_upload(update, context, dest, name)
        return
    if _tm_loaded():
        from ticketmaster.tm_worker_bridge import probe_tm_batch_file, run_tm_batch_import

        if force_tm or probe_tm_batch_file(dest):
            await run_tm_batch_import(update, context, dest, name)
            return

    await stock_upload(update, context, local_path=dest, local_name=name)


async def fulfill_stock_upload(
    update: Update,
    context: ContextTypes.DEFAULT_TYPE,
    path: Path,
    name: str,
) -> None:
    from vl_fulfill_stock_db import stats as ff_stats, upsert_rows
    from vl_fulfill_stock_parse import parse_fulfill_upload

    msg = await update.message.reply_text(f"Merging fulfill stock from {name}…")
    try:
        rows = await asyncio.to_thread(parse_fulfill_upload, path)
        if not rows:
            await msg.edit_text(
                "No barcoded seats found. Use tab-separated checker export "
                "(kind, email, password, event, date, venue, section, row, seat, barcode…)."
            )
            return
        result = await asyncio.to_thread(upsert_rows, rows, source_file=name)
        st = await asyncio.to_thread(ff_stats)
    except Exception as exc:
        log.exception("fulfill_file upload failed")
        await msg.edit_text(f"Fulfill file import failed: {exc}"[:4000])
        return
    await msg.edit_text(
        f"Fulfill stock merged from <b>{html.escape(name)}</b>\n"
        f"Parsed barcoded seats: <b>{len(rows)}</b>\n"
        f"New: <b>{result['inserted']}</b> · Updated: <b>{result['updated']}</b> · "
        f"Skipped: <b>{result['skipped']}</b>\n"
        f"Fulfill DB total: <b>{st['with_barcode']}</b> seats\n\n"
        "<code>/autofulfill</code> uses this file <b>first</b> for barcode matching.",
        parse_mode=ParseMode.HTML,
    )


async def stock_upload(
    update: Update,
    context: ContextTypes.DEFAULT_TYPE,
    *,
    local_path: Path | None = None,
    local_name: str | None = None,
) -> None:
    if not update.effective_user or not update.message:
        return
    uid = update.effective_user.id
    if not _admin(uid):
        if _worker(uid):
            await update.message.reply_text(_WORKER_ONLY_HINT)
        return
    if local_path is not None:
        dest = local_path
        name = local_name or dest.name
    else:
        if not update.message.document:
            return
        doc = update.message.document
        name = doc.file_name or "upload.txt"
        if not name.lower().endswith((".txt", ".csv", ".tsv", ".xlsx", ".xlsm")):
            return
        dest_dir = Path(tempfile.mkdtemp(prefix="vl_stock_in_"))
        dest = dest_dir / name
        tg_file = await doc.get_file()
        await tg_file.download_to_drive(custom_path=str(dest))
    if not name.lower().endswith((".txt", ".csv", ".tsv", ".xlsx", ".xlsm")):
        return
    msg = await update.message.reply_text(f"Importing ticket stock: {name}…")
    try:
        stats = await asyncio.to_thread(_import_stock_file, str(dest), name)
    except Exception as exc:
        await msg.edit_text(f"Import failed: {exc}")
        return
    if not stats["parsed"]:
        hint = ""
        if _tm_loaded() and name.lower().endswith((".xlsx", ".xls", ".txt")):
            hint = "\n\nIf this is a <b>/work</b> account batch, use /upload then send the file again."
        await msg.edit_text(
            "No tickets recognized in that file.\n"
            "Supported: app-event lines, barcode lines, transfer links, "
            "Credential/Seat blocks, and the inventory CSV."
            + hint,
            parse_mode=ParseMode.HTML,
        )
        return
    kind_bits = ", ".join(f"{k} {n}" for k, n in sorted(stats["kinds"].items()))
    await msg.edit_text(
        f"Imported {name}\n"
        f"Parsed: {stats['parsed']} ({kind_bits})\n"
        f"New: {stats['inserted']} · Updated: {stats['updated']}\n\n"
        + _stock_summary_text()
    )


_WORKER_BOT_COMMANDS = [
    BotCommand("start", "Help"),
    BotCommand("list", "Consign a Victory Live listing"),
    BotCommand("work", "TM account listing job"),
    BotCommand("status", "TM job progress"),
    BotCommand("balance", "VL + TM earnings"),
    BotCommand("leaderboard", "Top workers"),
]

_ADMIN_BOT_COMMANDS = [
    BotCommand("sync", "Sync inventory and orders, then send files"),
    BotCommand("autofulfill", "Fulfill open VL orders from DB barcodes"),
    BotCommand("cancel_autofulfill", "Stop auto-fulfill and send partial report"),
    BotCommand("autofulfill_reset", "Clear autofulfill tried history"),
    BotCommand("fulfill_file", "Upload barcode file for autofulfill"),
    BotCommand("summary", "Database inventory totals"),
    BotCommand("listed", "Txt and Excel of listed and unlisted"),
    BotCommand("bydate", "Txt and Excel of inventory filtered by event date"),
    BotCommand("stockfrom", "Non-SF tickets: start + end dates, then filters"),
    BotCommand("dbstatus", "VL + TM + SF database summary"),
    BotCommand("stockdedupe", "Dedupe ticket stock DB"),
    BotCommand("upload", "TM worker batch (.xlsx/.txt) then /work"),
    BotCommand("batches", "TM batch queue stats"),
    BotCommand("collect", "Collect TM worker payouts"),
    BotCommand("stats", "Database totals and live Victory Live totals"),
    BotCommand("fetchibrahim", "Ibrahim lister DB + VL stats"),
    BotCommand("list", "Consign a listing"),
    BotCommand("balance", "Your listing balance"),
    BotCommand("stock", "Stored ticket totals"),
    BotCommand("stockfind", "Search stored tickets"),
    BotCommand("stockfile", "Export every stored ticket"),
    BotCommand("whitelist", "Add, remove, or list users"),
    BotCommand("blacklist", "Block or unblock users"),
    BotCommand("addbal", "Add or remove a user balance"),
]


async def post_init(application: Application) -> None:
    await application.bot.set_my_commands(
        _WORKER_BOT_COMMANDS,
        scope=BotCommandScopeAllPrivateChats(),
    )
    for admin_id in sorted(ADMIN_IDS):
        try:
            await application.bot.set_my_commands(
                _ADMIN_BOT_COMMANDS,
                scope=BotCommandScopeChat(chat_id=admin_id),
            )
        except Exception as exc:
            log.warning("setMyCommands for admin %s failed: %s", admin_id, exc)

    from vl_auth import proxy_status_for_log
    from vl_inventory_sync import load_session

    log.info("Victory Live egress: %s (use_proxy=%s)", proxy_status_for_log(), USE_PROXY)

    async def _warm_vl_egress() -> None:
        try:
            from vl_auth import verify_egress_public_ip

            s = await asyncio.to_thread(load_session, USE_PROXY)
            pub_ip = await asyncio.to_thread(verify_egress_public_ip, s)
            log.info("Victory Live egress public IP: %s (VL calls must show 96.x)", pub_ip)
            code = await asyncio.to_thread(
                lambda: s.get(
                    "https://one.victorylive.com/api/security/profile",
                    timeout=45,
                    allow_redirects=True,
                ).status_code
            )
            log.info("Victory Live startup probe via egress: HTTP %s", code)
        except Exception as exc:
            log.warning("Victory Live startup probe failed (egress/proxy/session): %s", exc)

    await _warm_vl_egress()

    from vl_rm_recovery import backfill_pending_from_log, notify_recovered, run_pending_retries

    bf = await asyncio.to_thread(backfill_pending_from_log)
    log.info("rm pending backfill from log: %s", bf)

    async def rm_pending_job(context: ContextTypes.DEFAULT_TYPE) -> None:
        result = await asyncio.to_thread(
            run_pending_retries,
            affiliate_group_id=AFFILIATE_GROUP_ID,
            use_proxy=USE_PROXY,
            listing_payout_usd=LISTING_PAYOUT_USD,
            limit=10,
        )
        payloads = result.get("payloads") or []
        if payloads:
            await notify_recovered(context.bot, payloads, set(ADMIN_IDS))
        if result.get("ok") or result.get("fail"):
            log.info(
                "rm pending retry ok=%s fail=%s",
                result.get("ok"),
                result.get("fail"),
            )

    if application.job_queue:
        application.job_queue.run_repeating(rm_pending_job, interval=180, first=45)
    else:
        log.warning("job_queue unavailable — rm pending retries disabled")


def main() -> None:
    if not BOT_TOKEN:
        raise SystemExit("Set VL_BOT_TOKEN or vl_inventory_bot.config.json bot_token")
    db.init_db()
    app = (
        Application.builder()
        .token(BOT_TOKEN)
        .concurrent_updates(True)
        .post_init(post_init)
        .build()
    )
    listing_conv, listing_text, listing_callbacks, listing_cancel = build_listing_conversation(
        can_use=_can_list,
        use_proxy=USE_PROXY,
        client_po_profile_id=CLIENT_PO_PROFILE_ID,
        profile_name=CONSIGN_PROFILE_NAME,
        affiliate_group_id=AFFILIATE_GROUP_ID,
        admin_ids=set(ADMIN_IDS),
        listing_payout_usd=LISTING_PAYOUT_USD,
    )

    async def unified_cancel_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
        if context.user_data.get("stockfrom_await"):
            _stockfrom_clear_await(context, update.effective_user.id if update.effective_user else None)
            if update.message:
                await update.message.reply_text("Stock export cancelled.")
            raise ApplicationHandlerStop()
        if context.user_data.get("vl_list_await") or context.user_data.get("vl_list"):
            await listing_cancel(update, context)
            raise ApplicationHandlerStop()
        try:
            from ticketmaster import bot as tm

            await tm.cancel_cmd(update, context)
        except Exception:
            if update.message:
                await update.message.reply_text("Nothing to cancel.")

    # /stockfrom first (group -1) so dates never hit /list or TM routers.
    app.add_handler(CommandHandler("stockfrom", stockfrom_cmd), group=-1)
    app.add_handler(
        MessageHandler(filters.TEXT & ~filters.COMMAND, stockfrom_date_reply, block=True),
        group=-1,
    )
    app.add_handler(listing_callbacks)
    app.add_handler(listing_text)
    app.add_handler(listing_conv)
    app.add_handler(CommandHandler("cancel", unified_cancel_cmd))
    app.add_handler(CommandHandler("start", start_cmd))
    app.add_handler(CommandHandler("help", start_cmd))
    app.add_handler(CommandHandler("sync", sync_cmd))
    app.add_handler(CommandHandler("autofulfill_reset", autofulfill_reset_cmd))
    app.add_handler(CommandHandler("cancel_autofulfill", cancel_autofulfill_cmd))
    app.add_handler(CommandHandler("autofulfill", autofulfill_cmd))
    app.add_handler(CommandHandler("fulfill_file", fulfill_file_cmd))
    app.add_handler(CommandHandler("fixstocktype", fixstocktype_cmd))
    app.add_handler(CommandHandler("summary", summary_cmd))
    app.add_handler(CommandHandler("listed", listed_cmd))
    app.add_handler(CommandHandler("bydate", bydate_cmd))
    app.add_handler(CommandHandler("stats", stats_cmd))
    app.add_handler(CommandHandler("fetchibrahim", fetchibrahim_cmd))
    app.add_handler(CommandHandler("whitelist", whitelist_cmd))
    app.add_handler(CommandHandler("blacklist", blacklist_cmd))
    app.add_handler(CommandHandler("addbal", addbal_cmd))
    app.add_handler(CommandHandler("balance", balance_cmd))
    app.add_handler(CommandHandler("stock", stock_cmd))
    app.add_handler(CommandHandler("stockfind", stockfind_cmd))
    app.add_handler(CommandHandler("stockfile", stockfile_cmd))
    app.add_handler(CommandHandler("dbstatus", dbstatus_cmd))
    app.add_handler(CommandHandler("stockdedupe", stockdedupe_cmd))
    app.add_handler(CommandHandler("upload", upload_batch_cmd))
    app.add_handler(CallbackQueryHandler(stockfrom_button, pattern=r"^stf:"))
    app.add_handler(MessageHandler(filters.Document.ALL, admin_document_router))
    try:
        from ticketmaster.tm_worker_bridge import attach_tm_worker_handlers

        attach_tm_worker_handlers(
            app,
            repo_root=ROOT,
            admin_ids=ADMIN_IDS,
            is_admin=_admin,
            is_worker=_worker,
        )
        from ticketmaster.tm_worker_bridge import sync_worker_to_tm_whitelist

        for row in db.list_workers():
            sync_worker_to_tm_whitelist(int(row["telegram_id"]))
        log.info("TM worker module loaded")
    except Exception as exc:
        log.warning("TM worker module not loaded: %s", exc)
    async def telegram_error_handler(update: object, context: ContextTypes.DEFAULT_TYPE) -> None:
        log.error("Handler error", exc_info=context.error)
        if isinstance(update, Update) and update.effective_message:
            try:
                await update.effective_message.reply_text(
                    "Something went wrong — try again or /cancel.\n"
                    f"<code>{html.escape(str(context.error)[:300])}</code>",
                    parse_mode=ParseMode.HTML,
                )
            except Exception:
                pass

    app.add_error_handler(telegram_error_handler)
    log.info("VL inventory bot starting (admins=%s)", sorted(ADMIN_IDS))
    app.run_polling(allowed_updates=Update.ALL_TYPES)


if __name__ == "__main__":
    main()
